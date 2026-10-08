#!/usr/bin/env python3
"""Deterministic Canon E6 corpus materializer and benchmark lifecycle harness.

This tool never emits an SLO verdict. It materializes a pinned synthetic baseline
corpus and can bind an ordinary unclassified production trace to benchmark
lifecycle evidence only after actually running predeclared reset/prime and
measurement commands without a shell.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import subprocess
import time
from typing import Any
import zipfile

from docx import Document

from scripts import performance_slo_gate as gate
from scripts import performance_trace_observation as trace_obs

CORPUS_SCHEMA = "dokkomplekt.performance-corpus-spec.v1"
CORPUS_MANIFEST_SCHEMA = "dokkomplekt.performance-corpus-manifest.v1"
COMMAND_SCHEMA = "dokkomplekt.performance-benchmark-command.v1"
CONTEXT_SCHEMA = "dokkomplekt.performance-benchmark-context.v1"
PRODUCER = "performance_benchmark_harness.py"
CONTEXT_CLAIM = "benchmark_lifecycle_context_only_not_slo_verdict"
CORPUS_CLAIM = "deterministic_synthetic_baseline_not_production_measurement"
FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)

SPEC_PROFILE_KEYS = {
    "documents", "paragraphs", "paragraph_bytes", "tables", "rows", "cols",
    "header_footer", "roles",
}
CONTEXT_KEYS = {
    "schema", "producer", "claim", "session_id", "measurement_plan_sha256",
    "corpus_spec_sha256", "source_sha256", "trace_sha256", "cache_state",
    "run_phase", "cache_action", "cache_command_sha256",
    "measurement_command_sha256", "prior_trace_sha256", "cache_duration_ms",
    "measurement_duration_ms",
}


def _load(path: Path) -> dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError(f"{path}: root must be an object")
    return value


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _opaque(value: Any, label: str) -> str:
    if (
        not isinstance(value, str)
        or not value
        or len(value) > 128
        or not all(ch.isascii() and (ch.isalnum() or ch in "-_.") for ch in value)
    ):
        raise ValueError(f"{label} must be an opaque ASCII identifier")
    return value


def validate_corpus_spec(spec: dict[str, Any]) -> None:
    if set(spec) != {"schema", "corpus_id", "claim", "generator", "classes"}:
        raise ValueError("corpus spec keys must be closed")
    if spec.get("schema") != CORPUS_SCHEMA:
        raise ValueError(f"corpus spec schema must be {CORPUS_SCHEMA}")
    _opaque(spec.get("corpus_id"), "corpus_id")
    if spec.get("claim") != CORPUS_CLAIM:
        raise ValueError("corpus spec claim is invalid")
    if spec.get("generator") != PRODUCER:
        raise ValueError("corpus spec generator is invalid")
    classes = spec.get("classes")
    if not isinstance(classes, dict) or set(classes) != set(gate.CANONICAL_CORPUS_CLASSES):
        raise ValueError("corpus spec must define exactly the canonical corpus classes")
    for class_name, profile in classes.items():
        if not isinstance(profile, dict) or set(profile) != SPEC_PROFILE_KEYS:
            raise ValueError(f"{class_name}: profile keys must be closed")
        for key in ("documents", "paragraphs", "paragraph_bytes", "tables", "rows", "cols", "roles"):
            value = profile.get(key)
            if not isinstance(value, int) or isinstance(value, bool) or value < 0:
                raise ValueError(f"{class_name}.{key} must be a non-negative integer")
        if profile["documents"] <= 0 or profile["paragraph_bytes"] <= 0 or profile["roles"] <= 0:
            raise ValueError(f"{class_name}: documents/paragraph_bytes/roles must be positive")
        if not isinstance(profile.get("header_footer"), bool):
            raise ValueError(f"{class_name}.header_footer must be boolean")
    if classes["batch_10"]["documents"] != 10:
        raise ValueError("batch_10 must materialize exactly 10 documents")
    if classes["batch_50"]["documents"] != 50:
        raise ValueError("batch_50 must materialize exactly 50 documents")


def _deterministic_text(seed: str, size: int) -> str:
    chunks: list[str] = []
    counter = 0
    while sum(len(chunk) for chunk in chunks) < size:
        chunks.append(hashlib.sha256(f"{seed}:{counter}".encode()).hexdigest())
        counter += 1
    return "".join(chunks)[:size]


def _normalize_docx(path: Path) -> None:
    with zipfile.ZipFile(path, "r") as source:
        entries = [(info.filename, source.read(info.filename)) for info in source.infolist()]
    temp = path.with_suffix(".normalized.docx")
    with zipfile.ZipFile(temp, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as target:
        for name, payload in sorted(entries):
            info = zipfile.ZipInfo(name, FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target.writestr(info, payload)
    temp.replace(path)


def _materialize_document(path: Path, class_name: str, index: int, profile: dict[str, Any]) -> None:
    document = Document()
    document.add_heading(f"E6 {class_name} fixture {index:03d}", level=1)
    for role in range(profile["roles"]):
        document.add_paragraph(f"role_{role:03d}: " + _deterministic_text(f"{class_name}:role:{role}", 96))
    for paragraph in range(profile["paragraphs"]):
        paragraph_seed = (
            paragraph % 8 if class_name == "repeated_blocks" else paragraph
        )
        document.add_paragraph(
            _deterministic_text(
                f"{class_name}:{index}:paragraph:{paragraph_seed}",
                profile["paragraph_bytes"],
            )
        )
    for table_index in range(profile["tables"]):
        table = document.add_table(rows=profile["rows"], cols=profile["cols"])
        for row_index, row in enumerate(table.rows):
            for col_index, cell in enumerate(row.cells):
                cell.text = _deterministic_text(
                    (
                        f"{class_name}:{index}:table:"
                        f"{table_index % 2 if class_name == 'repeated_blocks' else table_index}:"
                        f"{row_index % 4 if class_name == 'repeated_blocks' else row_index}:"
                        f"{col_index}"
                    ),
                    48,
                )
    if profile["header_footer"]:
        section = document.sections[0]
        section.header.paragraphs[0].text = _deterministic_text(f"{class_name}:header", 96)
        section.footer.paragraphs[0].text = _deterministic_text(f"{class_name}:footer", 96)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    _normalize_docx(path)


def materialize_corpus(spec_path: Path, output_dir: Path) -> dict[str, Any]:
    spec = _load(spec_path)
    validate_corpus_spec(spec)
    files: list[dict[str, Any]] = []
    for class_name in gate.CANONICAL_CORPUS_CLASSES:
        profile = spec["classes"][class_name]
        for index in range(1, profile["documents"] + 1):
            path = output_dir / class_name / f"{class_name}-{index:03d}.docx"
            _materialize_document(path, class_name, index, profile)
            files.append(
                {
                    "class": class_name,
                    "index": index,
                    "sha256": _sha256(path),
                    "bytes": path.stat().st_size,
                    "relative_path": path.relative_to(output_dir).as_posix(),
                }
            )
    manifest = {
        "schema": CORPUS_MANIFEST_SCHEMA,
        "claim": CORPUS_CLAIM,
        "corpus_id": spec["corpus_id"],
        "corpus_spec_sha256": _sha256(spec_path),
        "files": files,
    }
    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / "corpus-manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    return manifest


def _load_command(path: Path) -> tuple[list[str], str]:
    document = _load(path)
    if set(document) != {"schema", "argv"} or document.get("schema") != COMMAND_SCHEMA:
        raise ValueError(f"{path}: invalid benchmark command schema")
    argv = document.get("argv")
    if (
        not isinstance(argv, list)
        or not argv
        or any(not isinstance(item, str) or not item for item in argv)
    ):
        raise ValueError(f"{path}: argv must be a non-empty string array")
    return list(argv), _sha256(path)


def _expand(argv: list[str], source: Path, trace: Path) -> list[str]:
    return [
        token.replace("{source}", str(source)).replace("{trace}", str(trace))
        for token in argv
    ]


def _run(argv: list[str]) -> float:
    started = time.monotonic()
    completed = subprocess.run(argv, shell=False, check=False)
    elapsed_ms = (time.monotonic() - started) * 1000.0
    if completed.returncode != 0:
        raise ValueError(f"benchmark command failed with exit code {completed.returncode}")
    return elapsed_ms


def validate_context(document: dict[str, Any]) -> None:
    if set(document) != CONTEXT_KEYS:
        raise ValueError("benchmark context keys must be closed")
    if (
        document.get("schema") != CONTEXT_SCHEMA
        or document.get("producer") != PRODUCER
        or document.get("claim") != CONTEXT_CLAIM
    ):
        raise ValueError("benchmark context provenance is invalid")
    _opaque(document.get("session_id"), "benchmark session_id")
    for key in (
        "measurement_plan_sha256", "corpus_spec_sha256", "source_sha256",
        "trace_sha256", "cache_command_sha256", "measurement_command_sha256",
    ):
        value = document.get(key)
        if not isinstance(value, str) or len(value) != 64 or any(ch not in "0123456789abcdef" for ch in value):
            raise ValueError(f"benchmark context {key} is invalid")
    if document.get("cache_state") not in gate.CANONICAL_CACHE_STATES:
        raise ValueError("benchmark context cache_state is invalid")
    if document.get("run_phase") not in ("first_run", "repeat_run", "unclassified"):
        raise ValueError("benchmark context run_phase is invalid")
    cache_action = document.get("cache_action")
    if cache_action not in ("cold_reset", "warm_prime"):
        raise ValueError("benchmark context cache_action is invalid")
    expected_cache_action = (
        "cold_reset" if document["cache_state"] == "cold_cache" else "warm_prime"
    )
    if cache_action != expected_cache_action:
        raise ValueError("benchmark context cache_action does not match cache_state")
    prior = document.get("prior_trace_sha256")
    if prior is not None and (
        not isinstance(prior, str) or len(prior) != 64 or any(ch not in "0123456789abcdef" for ch in prior)
    ):
        raise ValueError("benchmark context prior_trace_sha256 is invalid")
    run_phase = document["run_phase"]
    if run_phase == "first_run" and prior is not None:
        raise ValueError("first_run benchmark context cannot contain prior trace evidence")
    if run_phase == "repeat_run" and prior is None:
        raise ValueError("repeat_run benchmark context requires prior trace evidence")
    for key in ("cache_duration_ms", "measurement_duration_ms"):
        value = document.get(key)
        if not isinstance(value, (int, float)) or isinstance(value, bool) or float(value) < 0:
            raise ValueError(f"benchmark context {key} is invalid")


def run_benchmark(
    *,
    spec_path: Path,
    plan_path: Path,
    reference_path: Path,
    source_path: Path,
    trace_path: Path,
    cache_command_path: Path,
    measurement_command_path: Path,
    session_dir: Path,
    session_id: str,
) -> dict[str, Any]:
    spec = _load(spec_path)
    validate_corpus_spec(spec)
    plan = _load(plan_path)
    if plan.get("schema") != trace_obs.PLAN_SCHEMA:
        raise ValueError("benchmark harness requires a trace measurement plan")
    reference = _load(reference_path)
    reference_errors = gate.validate_reference(reference, gate.load_object(
        Path(__file__).resolve().parents[1] / "performance" / "slo-targets.json"
    ))
    if reference_errors:
        raise ValueError("invalid benchmark reference policy: " + "; ".join(reference_errors))
    if plan.get("reference_id") != reference.get("reference_id"):
        raise ValueError("measurement plan reference_id differs from benchmark reference")
    if plan.get("corpus_id") != reference.get("corpus_id"):
        raise ValueError("measurement plan corpus_id differs from benchmark reference")
    if spec.get("corpus_id") != plan.get("corpus_id"):
        raise ValueError("corpus spec corpus_id differs from measurement plan")
    expected_spec_sha256 = reference.get("corpus_spec_sha256")
    if expected_spec_sha256 != _sha256(spec_path):
        raise ValueError("corpus spec hash differs from benchmark reference")
    cache_state = plan.get("cache_state")
    if cache_state not in gate.CANONICAL_CACHE_STATES:
        raise ValueError("measurement plan cache_state is not benchmark-classifiable")
    run_kind = plan.get("run_kind")
    run_phase = run_kind if run_kind in ("first_run", "repeat_run") else "unclassified"
    expected_source = plan.get("expected_source_sha256")
    source_sha256 = _sha256(source_path)
    if expected_source != source_sha256:
        raise ValueError("benchmark source does not match plan.expected_source_sha256")
    session_id = _opaque(session_id, "session_id")
    state_path = session_dir / f"{session_id}.json"
    prior_trace_sha256: str | None = None
    if run_phase == "first_run":
        if state_path.exists():
            raise ValueError("first_run requires a fresh benchmark session")
    elif run_phase == "repeat_run":
        if not state_path.is_file():
            raise ValueError("repeat_run requires a completed prior measurement in this session")
        state = _load(state_path)
        if state.get("source_sha256") != source_sha256:
            raise ValueError("repeat_run prior measurement belongs to a different source")
        prior_trace_sha256 = state.get("trace_sha256")
        if not isinstance(prior_trace_sha256, str):
            raise ValueError("repeat_run prior trace evidence is invalid")

    cache_argv, cache_command_sha256 = _load_command(cache_command_path)
    measurement_argv, measurement_command_sha256 = _load_command(measurement_command_path)
    cache_duration_ms = _run(_expand(cache_argv, source_path, trace_path))
    if trace_path.exists():
        trace_path.unlink()
    measurement_duration_ms = _run(_expand(measurement_argv, source_path, trace_path))
    if not trace_path.is_file():
        raise ValueError("measurement command did not create the declared trace artifact")
    trace = _load(trace_path)
    trace_info = trace_obs._validate_trace(trace)
    if trace_info.get("source_sha256") != source_sha256:
        raise ValueError("measured trace source does not match benchmark source")
    trace_sha256 = _sha256(trace_path)
    context = {
        "schema": CONTEXT_SCHEMA,
        "producer": PRODUCER,
        "claim": CONTEXT_CLAIM,
        "session_id": session_id,
        "measurement_plan_sha256": _sha256(plan_path),
        "corpus_spec_sha256": _sha256(spec_path),
        "source_sha256": source_sha256,
        "trace_sha256": trace_sha256,
        "cache_state": cache_state,
        "run_phase": run_phase,
        "cache_action": "cold_reset" if cache_state == "cold_cache" else "warm_prime",
        "cache_command_sha256": cache_command_sha256,
        "measurement_command_sha256": measurement_command_sha256,
        "prior_trace_sha256": prior_trace_sha256,
        "cache_duration_ms": cache_duration_ms,
        "measurement_duration_ms": measurement_duration_ms,
    }
    validate_context(context)
    session_dir.mkdir(parents=True, exist_ok=True)
    state_path.write_text(
        json.dumps(
            {
                "source_sha256": source_sha256,
                "trace_sha256": trace_sha256,
                "measurement_command_sha256": measurement_command_sha256,
            },
            sort_keys=True,
        ) + "\n",
        encoding="utf-8",
    )
    return context


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Materialize E6 corpus or run benchmark lifecycle")
    sub = parser.add_subparsers(dest="command", required=True)

    materialize = sub.add_parser("materialize")
    materialize.add_argument("--spec", default="performance/corpus-spec.json", type=Path)
    materialize.add_argument("--output-dir", required=True, type=Path)

    run = sub.add_parser("run")
    run.add_argument("--spec", default="performance/corpus-spec.json", type=Path)
    run.add_argument("--plan", required=True, type=Path)
    run.add_argument("--reference", required=True, type=Path)
    run.add_argument("--source", required=True, type=Path)
    run.add_argument("--trace", required=True, type=Path)
    run.add_argument("--cache-command", required=True, type=Path)
    run.add_argument("--measurement-command", required=True, type=Path)
    run.add_argument("--session-dir", required=True, type=Path)
    run.add_argument("--session-id", required=True)
    run.add_argument("--output", required=True, type=Path)
    args = parser.parse_args(argv)

    try:
        if args.command == "materialize":
            manifest = materialize_corpus(args.spec, args.output_dir)
            print(f"PERFORMANCE CORPUS MATERIALIZED: files={len(manifest['files'])}; no SLO verdict")
            return 0
        context = run_benchmark(
            spec_path=args.spec,
            plan_path=args.plan,
            reference_path=args.reference,
            source_path=args.source,
            trace_path=args.trace,
            cache_command_path=args.cache_command,
            measurement_command_path=args.measurement_command,
            session_dir=args.session_dir,
            session_id=args.session_id,
        )
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(
            json.dumps(context, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
            encoding="utf-8",
        )
        print(
            "BENCHMARK CONTEXT OK: "
            f"cache={context['cache_state']} phase={context['run_phase']}; no SLO verdict"
        )
        return 0
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        print(f"PERFORMANCE BENCHMARK FAILED: {exc}")
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
